import os
import docx
import pandas as pd
from typing import Optional, Sequence

from data_processing import DataProcessing, MetricTable, TTestTable


class ExportDocument:
    @classmethod
    def export_metric_table_docx(
        cls, output_dir: str, metric_table: MetricTable
    ):
        mean_df = metric_table.mean
        std_df = metric_table.std

        if std_df is not None:
            table = (
                (mean_df * 100).applymap(lambda x: "{:.2f}".format(x))
                + u" \u00B1 "
                + (std_df * 100).applymap(lambda x: "{:.2f}".format(x))
            )
        else:
            table = (mean_df * 100).applymap(lambda x: "{:.2f}".format(x))

        final_mean = mean_df.mean(axis=0)

        doc = docx.Document()
        t = doc.add_table(table.shape[0] + 2, table.shape[1] + 1)
        for i in range(table.shape[0]):
            t.cell(i + 1, 0).text = table.index[i]
        for j in range(table.shape[-1]):
            t.cell(0, j + 1).text = table.columns[j]
        for i in range(table.shape[0]):
            for j in range(table.shape[-1]):
                t.cell(i + 1, j + 1).text = str(table.values[i, j])

        final_row = table.shape[0] + 1
        t.cell(final_row, 0).text = "Mean {}".format(
            metric_table.metric_name.capitalize()
        )
        for i, fm in enumerate(final_mean.values, start=1):
            t.cell(final_row, i).text = "{:.2f}".format(fm)

        output_dir = os.path.abspath(output_dir)
        if not os.path.isdir(output_dir):
            os.makedirs(output_dir)
        file_prefix = "_".join(metric_table.group.values())
        if not file_prefix:
            filename = "{}_table".format(metric_table.metric_name)
        else:
            filename = "{}_{}_table".format(
                file_prefix, metric_table.metric_name
            )
        output_path = os.path.join(output_dir, "{}.docx".format(filename),)
        doc.save(output_path)

    @classmethod
    def export_csv(
        cls,
        output_dir: str,
        dataset_name: str,
        grouped_df: pd.DataFrame,
        overall_df: pd.DataFrame,
        file_suffix: Optional[str] = None,
    ):
        if "dataset" in grouped_df.columns:
            assert (grouped_df["dataset"] == dataset_name).all()
        if "dataset" in overall_df.columns:
            assert (overall_df["dataset"] == dataset_name).all()
        output_dir = os.path.abspath(output_dir)
        if not os.path.isdir(output_dir):
            os.makedirs(output_dir)

        filename = (
            "{}_{}_results".format(dataset_name, file_suffix)
            if file_suffix
            else "{}_results".format(dataset_name)
        )
        output_path = os.path.join(output_dir, "{}.csv".format(filename))
        grouped_df.to_csv(
            output_path,
            index=False,
            columns=[col for col in grouped_df if "_list" not in col],
        )

        filename += "_overall"
        output_path = os.path.join(output_dir, "{}.csv".format(filename))
        overall_df.to_csv(
            output_path,
            index=False,
            columns=[col for col in overall_df if "_list" not in col],
        )

    @classmethod
    def export_excel_sheet(
        cls,
        output_dir: str,
        dataset_name: str,
        grouped_df: pd.DataFrame,
        overall_df: pd.DataFrame,
        grouped_ttest_table: Sequence[TTestTable],
        overall_ttest_table: Sequence[TTestTable],
        file_suffix: Optional[str] = None,
    ):
        if "dataset" in grouped_df.columns:
            assert (grouped_df["dataset"] == dataset_name).all()
        if "dataset" in overall_df.columns:
            assert (overall_df["dataset"] == dataset_name).all()

        for ttest_table in grouped_ttest_table:
            assert (
                ttest_table.group.get("dataset", dataset_name) == dataset_name
            )
        for ttest_table in overall_ttest_table:
            assert (
                ttest_table.group.get("dataset", dataset_name) == dataset_name
            )

        output_dir = os.path.abspath(output_dir)
        if not os.path.isdir(output_dir):
            os.makedirs(output_dir)
        filename = (
            "{}_{}_results".format(dataset_name, file_suffix)
            if file_suffix
            else "{}_results".format(dataset_name)
        )
        output_path = os.path.join(output_dir, "{}.xlsx".format(filename))

        with pd.ExcelWriter(output_path, "xlsxwriter") as writer:
            grouped_df.to_excel(
                writer,
                "{}_results".format(dataset_name),
                columns=[col for col in grouped_df if "_list" not in col],
                index=False,
                freeze_panes=None,
            )
            overall_df.to_excel(
                writer,
                "{}_overall_results".format(dataset_name),
                columns=[col for col in overall_df if "_list" not in col],
                index=False,
                freeze_panes=None,
            )

            ttest_df_dict = DataProcessing.merge_ttest_tables(
                grouped_ttest_table
            )
            for metric_name, ttest_df in ttest_df_dict.items():
                ttest_df.to_excel(
                    writer,
                    "{}_{}_ttest".format(dataset_name, metric_name),
                    freeze_panes=None,
                )

            ttest_df_dict = DataProcessing.merge_ttest_tables(
                overall_ttest_table
            )
            for metric_name, ttest_df in ttest_df_dict.items():
                ttest_df.to_excel(
                    writer,
                    "{}_{}_overall_ttest".format(dataset_name, metric_name),
                    freeze_panes=None,
                )

